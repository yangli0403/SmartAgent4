# -*- coding: utf-8 -*-
"""
SmartAgent4 项目分析报告 Word 文档生成脚本 - 含图表
"""

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Rectangle
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# 设置中文字体支持
plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'Arial Unicode MS', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False

def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(shading_elm)

def draw_architecture_diagram():
    """绘制系统架构图"""
    fig, ax = plt.subplots(1, 1, figsize=(14, 10))
    ax.set_xlim(0, 14)
    ax.set_ylim(0, 10)
    ax.axis('off')
    ax.set_title('SmartAgent4 系统架构图', fontsize=16, fontweight='bold', pad=20)
    
    # 颜色定义
    colors = {
        'client': '#E3F2FD',
        'gateway': '#E8F5E9',
        'service': '#FFF3E0',
        'datastore': '#F3E5F5',
        'external': '#FFEBEE',
        'border': '#333333'
    }
    
    # 绘制客户端层
    client_box = FancyBboxPatch((0.3, 7.5), 2.5, 2, boxstyle="round,pad=0.1", 
                                  facecolor=colors['client'], edgecolor='blue', linewidth=2)
    ax.add_patch(client_box)
    ax.text(1.55, 9.2, 'Client', fontsize=12, fontweight='bold', ha='center')
    ax.text(0.6, 8.5, 'Web App', fontsize=9, ha='center', bbox=dict(boxstyle='round', facecolor='white'))
    ax.text(2.5, 8.5, 'Cockpit\nPage', fontsize=9, ha='center', bbox=dict(boxstyle='round', facecolor='white'))
    ax.text(0.6, 7.7, 'Chat\nPage', fontsize=9, ha='center', bbox=dict(boxstyle='round', facecolor='white'))
    ax.text(2.5, 7.7, 'Memories\nPage', fontsize=9, ha='center', bbox=dict(boxstyle='round', facecolor='white'))
    
    # 绘制网关层
    gateway_box = FancyBboxPatch((3.5, 7.8), 2.2, 1.5, boxstyle="round,pad=0.1", 
                                  facecolor=colors['gateway'], edgecolor='green', linewidth=2)
    ax.add_patch(gateway_box)
    ax.text(4.6, 9.0, 'Gateway', fontsize=12, fontweight='bold', ha='center')
    ax.text(4.6, 8.3, 'tRPC Router\n+ WebSocket', fontsize=9, ha='center')
    
    # 绘制服务层
    service_box = FancyBboxPatch((6.5, 4.5), 4, 5, boxstyle="round,pad=0.1", 
                                  facecolor=colors['service'], edgecolor='orange', linewidth=2)
    ax.add_patch(service_box)
    ax.text(8.5, 9.2, 'Core Services', fontsize=12, fontweight='bold', ha='center')
    
    # 服务节点
    services = [
        ('Supervisor\nGraph', 7.2, 8.5),
        ('Classify\nNode', 7.2, 7.5),
        ('Plan\nNode', 9.8, 8.5),
        ('Execute\nNode', 9.8, 7.5),
        ('Respond\nNode', 7.2, 6.3),
        ('Memory\nSystem', 9.8, 6.3),
        ('Personality\nEngine', 7.2, 5.3),
        ('AIRI\nBridge', 9.8, 5.3),
    ]
    for name, x, y in services:
        rect = FancyBboxPatch((x-0.7, y-0.4), 1.4, 0.8, boxstyle="round,pad=0.05",
                               facecolor='white', edgecolor='gray', linewidth=1)
        ax.add_patch(rect)
        ax.text(x, y, name, fontsize=7, ha='center', va='center')
    
    # 绘制数据层
    datastore_box = FancyBboxPatch((11.2, 5.5), 2.3, 3.5, boxstyle="round,pad=0.1", 
                                    facecolor=colors['datastore'], edgecolor='purple', linewidth=2)
    ax.add_patch(datastore_box)
    ax.text(12.35, 8.7, 'Data Stores', fontsize=12, fontweight='bold', ha='center')
    ax.text(12.35, 7.8, 'PostgreSQL\n(Drizzle ORM)', fontsize=9, ha='center', 
            bbox=dict(boxstyle='round', facecolor='white', edgecolor='gray'))
    ax.text(12.35, 6.5, 'Vector DB\n(Embedding)', fontsize=9, ha='center',
            bbox=dict(boxstyle='round', facecolor='white', edgecolor='gray'))
    
    # 绘制外部服务
    external_box = FancyBboxPatch((0.3, 1), 4.5, 3, boxstyle="round,pad=0.1", 
                                    facecolor=colors['external'], edgecolor='red', linewidth=2)
    ax.add_patch(external_box)
    ax.text(2.55, 3.7, 'External Services', fontsize=12, fontweight='bold', ha='center')
    ax.text(0.8, 2.9, 'Manus API\n(gpt-4.1-mini)', fontsize=8, ha='center',
            bbox=dict(boxstyle='round', facecolor='white', edgecolor='gray'))
    ax.text(2.5, 2.9, 'Volcengine\nARK', fontsize=8, ha='center',
            bbox=dict(boxstyle='round', facecolor='white', edgecolor='gray'))
    ax.text(4.2, 2.9, 'Ali\nDashScope', fontsize=8, ha='center',
            bbox=dict(boxstyle='round', facecolor='white', edgecolor='gray'))
    ax.text(1.65, 1.7, 'AIRI Runtime\n(Live2D)', fontsize=8, ha='center',
            bbox=dict(boxstyle='round', facecolor='white', edgecolor='gray'))
    ax.text(3.55, 1.7, 'Emotions\nExpress', fontsize=8, ha='center',
            bbox=dict(boxstyle='round', facecolor='white', edgecolor='gray'))
    
    # 绘制连接线
    arrows = [
        # Client to Gateway
        ((2.8, 8.5), (3.5, 8.5), 'HTTPS'),
        ((2.8, 8.2), (3.5, 8.0), 'WebSocket'),
        # Gateway to Services
        ((5.7, 8.5), (6.5, 8.5), ''),
        # Services internal
        ((7.9, 8.5), (9.1, 8.5), ''),
        ((7.9, 7.5), (9.1, 7.5), ''),
        ((7.9, 6.5), (9.1, 6.5), ''),
        ((7.9, 5.5), (9.1, 5.5), ''),
        # Services to Datastore
        ((10.5, 7.5), (11.2, 7.8), 'SQL'),
        ((10.5, 6.5), (11.2, 6.5), 'Vector'),
        # Services to External
        ((7.5, 5.0), (3.5, 3.0), 'LLM'),
        ((9.8, 5.0), (3.5, 2.5), 'Render'),
        # External to AIRI
        ((3.0, 2.9), (3.0, 2.2), ''),
    ]
    
    for (x1, y1), (x2, y2), label in arrows:
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                   arrowprops=dict(arrowstyle='->', color='gray', lw=1.5))
    
    plt.tight_layout()
    img_path = 'D:/DEMO/SmartAgent4_demo/docs/architecture_diagram.png'
    os.makedirs(os.path.dirname(img_path), exist_ok=True)
    plt.savefig(img_path, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    return img_path


def draw_flow_diagram():
    """绘制对话处理流程图"""
    fig, ax = plt.subplots(1, 1, figsize=(12, 10))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 10)
    ax.axis('off')
    ax.set_title('SmartAgent4 对话处理流程', fontsize=16, fontweight='bold', pad=20)
    
    # 定义节点
    nodes = [
        (1, 9, '用户消息', '#E3F2FD', 'start'),
        (1, 7.5, 'contextEnrichNode\nPre-Retrieval + 混合检索', '#FFF3E0', 'process'),
        (1, 5.5, 'classifyNode\n意图分类', '#E8F5E9', 'process'),
        (1, 3.5, 'planNode\n任务规划', '#E3F2FD', 'process'),
        (5, 3.5, 'parallelExecute\nDAG并行执行', '#FFF3E0', 'process'),
        (9, 3.5, 'executeNode\n串行执行', '#FFF3E0', 'process'),
        (5, 1.5, 'memoryExtractionNode\n异步记忆提取', '#F3E5F5', 'process'),
        (9, 1.5, 'reflectionNode\n自进化反思', '#F3E5F5', 'process'),
        (7, -0.5, 'respondNode\n响应生成 + 情感标签', '#E8F5E9', 'process'),
    ]
    
    # 绘制节点
    for x, y, text, color, node_type in nodes:
        if node_type == 'start':
            shape = FancyBboxPatch((x-1, y-0.4), 2, 0.8, boxstyle="round,pad=0.1",
                                   facecolor=color, edgecolor='blue', linewidth=2)
        else:
            shape = FancyBboxPatch((x-1.2, y-0.5), 2.4, 1, boxstyle="round,pad=0.1",
                                   facecolor=color, edgecolor='gray', linewidth=1.5)
        ax.add_patch(shape)
        ax.text(x, y, text, fontsize=8, ha='center', va='center', fontweight='bold' if node_type == 'start' else 'normal')
    
    # 绘制连接线和箭头
    connections = [
        ((2, 8.6), (2, 8.0)),
        ((2, 7.0), (2, 6.0)),
        ((2, 5.0), (2, 4.0)),
        ((2, 3.0), (4.8, 3.5), 'DAG'),
        ((2, 3.0), (8.8, 3.5), '串行'),
        ((5, 2.8), (7, 0.0)),
        ((9, 2.8), (7, 0.0), '', 'left'),
        ((5, 1.0), (6.2, 0.0)),
        ((9, 1.0), (7.8, 0.0), '', 'left'),
        ((2, 4.5), (2, 5.0)),
    ]
    
    for conn in connections:
        if len(conn) == 2:
            (x1, y1), (x2, y2) = conn
            label = ''
            direction = ''
        elif len(conn) == 3:
            (x1, y1), (x2, y2), label = conn
            direction = ''
        else:
            (x1, y1), (x2, y2), label, direction = conn
        
        style = '->'
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                   arrowprops=dict(arrowstyle=style, color='#333', lw=1.5))
        if label:
            mid_x = (x1 + x2) / 2
            mid_y = (y1 + y2) / 2
            offset = 0.15 if direction != 'left' else -0.15
            ax.text(mid_x, mid_y + offset, label, fontsize=7, ha='center' if direction != 'left' else 'right', 
                    color='blue', fontweight='bold')
    
    # 添加图例
    legend_elements = [
        mpatches.Patch(facecolor='#E3F2FD', edgecolor='blue', label='开始/用户输入'),
        mpatches.Patch(facecolor='#FFF3E0', edgecolor='gray', label='处理节点'),
        mpatches.Patch(facecolor='#E8F5E9', edgecolor='gray', label='输出节点'),
        mpatches.Patch(facecolor='#F3E5F5', edgecolor='gray', label='异步节点'),
    ]
    ax.legend(handles=legend_elements, loc='lower right', fontsize=8)
    
    plt.tight_layout()
    img_path = 'D:/DEMO/SmartAgent4_demo/docs/flow_diagram.png'
    plt.savefig(img_path, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    return img_path


def draw_agent_diagram():
    """绘制多智能体协同架构图"""
    fig, ax = plt.subplots(1, 1, figsize=(12, 8))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8)
    ax.axis('off')
    ax.set_title('SmartAgent4 多智能体协同架构', fontsize=16, fontweight='bold', pad=20)
    
    # 绘制 Supervisor
    supervisor = FancyBboxPatch((4.5, 5.5), 3, 1.5, boxstyle="round,pad=0.1",
                                facecolor='#FFE0B2', edgecolor='#E65100', linewidth=2)
    ax.add_patch(supervisor)
    ax.text(6, 6.25, 'Supervisor', fontsize=12, fontweight='bold', ha='center')
    ax.text(6, 5.85, '(编排协调器)', fontsize=9, ha='center')
    
    # 绘制 Agent Card Registry
    registry = FancyBboxPatch((0.3, 2), 2.5, 2, boxstyle="round,pad=0.1",
                               facecolor='#E3F2FD', edgecolor='blue', linewidth=2)
    ax.add_patch(registry)
    ax.text(1.55, 3.7, 'Agent Card Registry', fontsize=10, fontweight='bold', ha='center')
    ax.text(1.55, 3.1, '动态发现\nJSON配置加载', fontsize=8, ha='center')
    ax.text(0.7, 2.5, 'fileAgent.json', fontsize=7, ha='center',
            bbox=dict(boxstyle='round', facecolor='white'))
    ax.text(2.3, 2.5, 'navigationAgent', fontsize=7, ha='center',
            bbox=dict(boxstyle='round', facecolor='white'))
    
    # 绘制 Parallel Execute Engine
    engine = FancyBboxPatch((4.5, 2), 3, 2, boxstyle="round,pad=0.1",
                             facecolor='#C8E6C9', edgecolor='green', linewidth=2)
    ax.add_patch(engine)
    ax.text(6, 3.7, 'Parallel Execute Engine', fontsize=10, fontweight='bold', ha='center')
    ax.text(6, 3.1, '(DAG 拓扑排序)', fontsize=8, ha='center')
    ax.text(5, 2.5, '批次1\n并行执行', fontsize=7, ha='center',
            bbox=dict(boxstyle='round', facecolor='white'))
    ax.text(7, 2.5, '批次2\n等待完成', fontsize=7, ha='center',
            bbox=dict(boxstyle='round', facecolor='white'))
    
    # 绘制 Domain Agents
    agents = [
        ('0.3', '0.5', 'General\nAgent', '通用对话', '#FFCCBC'),
        ('3.3', '0.5', 'File\nAgent', '文件操作\n(15工具)', '#B3E5FC'),
        ('6.3', '0.5', 'Navigation\nAgent', '导航出行\n(19工具)', '#C5CAE9'),
        ('9.3', '0.5', 'Multimedia\nAgent', '多媒体\n(8工具)', '#F8BBD9'),
    ]
    
    for x, y, name, desc, color in agents:
        x, y = float(x), float(y)
        agent = FancyBboxPatch((x, y), 2.2, 1.5, boxstyle="round,pad=0.1",
                                facecolor=color, edgecolor='gray', linewidth=1.5)
        ax.add_patch(agent)
        ax.text(x+1.1, y+1.0, name, fontsize=9, fontweight='bold', ha='center')
        ax.text(x+1.1, y+0.5, desc, fontsize=7, ha='center')
    
    # 绘制连接线
    # Supervisor -> Registry
    ax.annotate('', xy=(2.8, 3), xytext=(4.5, 5.5),
               arrowprops=dict(arrowstyle='->', color='blue', lw=1.5))
    ax.text(3.5, 4.5, '查询能力', fontsize=8, color='blue')
    
    # Registry -> Engine
    ax.annotate('', xy=(4.5, 3), xytext=(2.8, 2.5),
               arrowprops=dict(arrowstyle='->', color='blue', lw=1.5))
    ax.text(3.5, 2.2, '匹配结果', fontsize=8, color='blue')
    
    # Engine -> Agents
    for i, (x, y, _, _, _) in enumerate(agents):
        x = float(x) + 1.1
        y = 2
        ax.annotate('', xy=(x, 0.5), xytext=(y, 2),
                   arrowprops=dict(arrowstyle='->', color='green', lw=1, ls='--'))
    
    ax.text(6, 1.2, 'DAG 并行执行', fontsize=10, ha='center', 
            color='green', fontweight='bold')
    
    # 委托协议标注
    ax.annotate('', xy=(10.5, 1.5), xytext=(9.5, 1.5),
               arrowprops=dict(arrowstyle='->', color='red', lw=1.5))
    ax.text(11.5, 1.5, 'delegate()\n委托', fontsize=8, color='red', va='center')
    
    plt.tight_layout()
    img_path = 'D:/DEMO/SmartAgent4_demo/docs/agent_diagram.png'
    plt.savefig(img_path, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    return img_path


def draw_memory_diagram():
    """绘制三层记忆系统架构图"""
    fig, ax = plt.subplots(1, 1, figsize=(12, 8))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8)
    ax.axis('off')
    ax.set_title('SmartAgent4 三层记忆系统架构', fontsize=16, fontweight='bold', pad=20)
    
    # 绘制三层记忆
    memories = [
        (6, 6.5, '情景记忆\n(Episodic)', '对话历史、用户事件', '#E3F2FD', 'blue'),
        (6, 4, '语义记忆\n(Semantic)', '知识概念、事实信息', '#E8F5E9', 'green'),
        (6, 1.5, '人格记忆\n(Identity)', '人格配置、对话风格', '#FFF3E0', 'orange'),
    ]
    
    for x, y, name, desc, color, edge_color in memories:
        box = FancyBboxPatch((x-2.5, y-0.8), 5, 1.6, boxstyle="round,pad=0.1",
                             facecolor=color, edgecolor=edge_color, linewidth=2)
        ax.add_patch(box)
        ax.text(x, y+0.3, name, fontsize=12, fontweight='bold', ha='center')
        ax.text(x, y-0.3, desc, fontsize=9, ha='center', color='gray')
    
    # 绘制处理组件
    components = [
        (1, 6.5, 'Pre-Retrieval\nDecision', '#FCE4EC'),
        (11, 6.5, 'Hybrid\nSearch', '#E1F5FE'),
        (1, 4, 'Embedding\nService', '#FFF8E1'),
        (11, 4, 'Behavior\nDetector', '#E0F7FA'),
        (1, 1.5, 'Confidence\nEvolution', '#F1F8E9'),
        (11, 1.5, 'Forgetting\nService', '#FBE9E7'),
    ]
    
    for x, y, name, color in components:
        box = FancyBboxPatch((x-1.2, y-0.5), 2.4, 1, boxstyle="round,pad=0.05",
                             facecolor=color, edgecolor='gray', linewidth=1)
        ax.add_patch(box)
        ax.text(x, y, name, fontsize=8, ha='center', va='center')
    
    # 绘制连接线
    # 左侧组件到情景记忆
    ax.annotate('', xy=(3.5, 6.5), xytext=(2.2, 6.5),
               arrowprops=dict(arrowstyle='->', color='purple', lw=1.5))
    ax.text(2.8, 6.9, '决策', fontsize=7, color='purple')
    
    # 情景记忆到右侧组件
    ax.annotate('', xy=(9.8, 6.5), xytext=(8.5, 6.5),
               arrowprops=dict(arrowstyle='->', color='purple', lw=1.5))
    ax.text(9.1, 6.9, '检索', fontsize=7, color='purple')
    
    # 向下箭头表示层级关系
    ax.annotate('', xy=(6, 4.8), xytext=(6, 5.7),
               arrowprops=dict(arrowstyle='->', color='gray', lw=1.5, ls='--'))
    ax.annotate('', xy=(6, 2.3), xytext=(6, 3.2),
               arrowprops=dict(arrowstyle='->', color='gray', lw=1.5, ls='--'))
    
    # PostgreSQL 数据库
    db = FancyBboxPatch((4.5, -0.5), 3, 1, boxstyle="round,pad=0.1",
                        facecolor='#F3E5F5', edgecolor='purple', linewidth=2)
    ax.add_patch(db)
    ax.text(6, 0, 'PostgreSQL 16\n(Drizzle ORM)', fontsize=9, ha='center', fontweight='bold')
    
    plt.tight_layout()
    img_path = 'D:/DEMO/SmartAgent4_demo/docs/memory_diagram.png'
    plt.savefig(img_path, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    return img_path


def create_document():
    """创建完整的 Word 文档"""
    doc = Document()
    
    # ==================== 文档标题 ====================
    title = doc.add_heading('SmartAgent4 项目分析报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('智能对话交互系统架构与代码分析')
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_paragraph()
    
    # 生成图表
    print("正在生成系统架构图...")
    arch_img = draw_architecture_diagram()
    
    print("正在生成对话处理流程图...")
    flow_img = draw_flow_diagram()
    
    print("正在生成多智能体协同架构图...")
    agent_img = draw_agent_diagram()
    
    print("正在生成三层记忆系统图...")
    memory_img = draw_memory_diagram()
    
    # 插入架构图
    doc.add_heading('1. 系统架构图', 1)
    doc.add_picture(arch_img, width=Inches(6.5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # ==================== 2. 代码量统计 ====================
    doc.add_heading('2. 代码量统计', 1)
    doc.add_paragraph('项目代码总量统计：')
    
    table = doc.add_table(rows=4, cols=4)
    table.style = 'Table Grid'
    headers = ['层级', '文件数', '代码行数', '占比']
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True
        header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(header_cells[i], '4472C4')
    
    data = [
        ['前端 (Client)', '123', '17,514', '31.7%'],
        ['后端 (Server)', '192', '37,712', '68.3%'],
        ['总计', '315', '55,226', '100%'],
    ]
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_data in enumerate(row_data):
            row_cells[col_idx].text = cell_data
    
    doc.add_paragraph()
    
    # ==================== 3. 对话处理流程图 ====================
    doc.add_heading('3. 对话处理流程图', 1)
    doc.add_picture(flow_img, width=Inches(6))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('对话处理流程说明：')
    flow_steps = [
        '1. 用户消息输入',
        '2. contextEnrichNode - Pre-Retrieval Decision → 混合检索 + 画像构建',
        '3. classifyNode - 意图分类（动态 Prompt 注入 Agent 能力描述）',
        '4. planNode - 任务规划（动态 Agent 列表 + 并行执行提示）',
        '5. parallelExecute/executeNode - DAG并行执行 或 串行执行',
        '6. respondNode - 响应生成（含情感标签）',
        '7. memoryExtractionNode - 异步记忆提取 + 行为检测',
        '8. reflectionNode - 异步自进化反思',
        '9. AI 回复 + 情感渲染',
    ]
    for step in flow_steps:
        p = doc.add_paragraph(step)
        p.paragraph_format.left_indent = Inches(0.3)
    
    # ==================== 4. 多智能体协同架构图 ====================
    doc.add_heading('4. 多智能体协同架构图', 1)
    doc.add_picture(agent_img, width=Inches(6))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('多智能体协同核心机制：')
    mechanisms = [
        '• Agent Card 动态发现：新增 Agent 只需在 agent-cards/ 目录放置 JSON 配置文件',
        '• DAG 并行执行引擎：基于 PlanStep.dependsOn 构建有向无环图，拓扑排序并行执行',
        '• 委托协议：Domain Agent 可通过 delegate() 横向委托其他 Agent，深度限制3层',
    ]
    for m in mechanisms:
        p = doc.add_paragraph(m)
        p.paragraph_format.left_indent = Inches(0.3)
    
    # ==================== 5. 三层记忆系统架构图 ====================
    doc.add_heading('5. 三层记忆系统架构图', 1)
    doc.add_picture(memory_img, width=Inches(6))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('三层记忆系统说明：')
    memory_desc = [
        '• 情景记忆 (Episodic)：对话历史、用户事件记录',
        '• 语义记忆 (Semantic)：知识概念、事实信息存储',
        '• 人格记忆 (Identity)：人格配置、对话风格参数',
    ]
    for d in memory_desc:
        p = doc.add_paragraph(d)
        p.paragraph_format.left_indent = Inches(0.3)
    
    doc.add_paragraph()
    
    # ==================== 6. 技术栈 ====================
    doc.add_heading('6. 技术栈', 1)
    
    table2 = doc.add_table(rows=14, cols=2)
    table2.style = 'Table Grid'
    tech_data = [
        ['层级', '技术选型'],
        ['前端框架', 'React 19 + TypeScript + Vite 7'],
        ['UI 库', 'Radix UI + TailwindCSS 4'],
        ['状态管理', 'Zustand + TanStack Query'],
        ['2D渲染', 'PixiJS 7 + pixi-live2d-display'],
        ['后端框架', 'Node.js + Express + tRPC 11'],
        ['AI框架', 'LangGraph (StateGraph) + LangChain'],
        ['LLM', 'Manus API + Volcengine ARK (双轨)'],
        ['向量检索', '阿里云百炼 DashScope (1024维)'],
        ['数据库', 'PostgreSQL 16 + Drizzle ORM'],
        ['实时通信', 'WebSocket (ws)'],
        ['协议', 'MCP (Model Context Protocol)'],
        ['测试', 'Vitest 2 + @vitest/coverage-v8'],
        ['包管理', 'pnpm 10'],
    ]
    for row_idx, row_data in enumerate(tech_data):
        row_cells = table2.rows[row_idx].cells
        for col_idx, cell_data in enumerate(row_data):
            row_cells[col_idx].text = cell_data
            if row_idx == 0:
                row_cells[col_idx].paragraphs[0].runs[0].bold = True
                row_cells[col_idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(row_cells[col_idx], '4472C4')
    
    doc.add_paragraph()
    
    # ==================== 7. 核心功能模块 ====================
    doc.add_heading('7. 核心功能模块', 1)
    
    table3 = doc.add_table(rows=9, cols=3)
    table3.style = 'Table Grid'
    modules = [
        ['模块', '描述', '文件数'],
        ['Agent 编排', 'LangGraph Supervisor 多节点编排', '~30'],
        ['多智能体协同', 'Agent Card 发现 + DAG并行执行', '~15'],
        ['三层记忆系统', '情景/语义/人格记忆 + 遗忘机制', '~25'],
        ['个性引擎', '人格配置加载与动态 Prompt 组装', '~5'],
        ['情感表达', '情感标签渲染 + AIRI 桥接', '~8'],
        ['AIRI 舞台', 'Live2D 前端渲染层', '~15'],
        ['MCP 工具', '网易云音乐、文件管理、导航等', '~20'],
        ['UI 组件库', 'Radix UI 封装组件', '51'],
    ]
    for row_idx, row_data in enumerate(modules):
        row_cells = table3.rows[row_idx].cells
        for col_idx, cell_data in enumerate(row_data):
            row_cells[col_idx].text = cell_data
            if row_idx == 0:
                row_cells[col_idx].paragraphs[0].runs[0].bold = True
                row_cells[col_idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(row_cells[col_idx], '4472C4')
    
    doc.add_paragraph()
    
    # ==================== 8. 测试覆盖 ====================
    doc.add_heading('8. 测试覆盖', 1)
    
    table4 = doc.add_table(rows=8, cols=4)
    table4.style = 'Table Grid'
    tests = [
        ['模块', '语句覆盖率', '函数覆盖率', '测试用例'],
        ['AIRI 舞台核心库', '100%', '100%', '71'],
        ['discovery 模块', '97.68%', '100%', '77'],
        ['embeddingService', '98.4%', '100%', '24'],
        ['extractionAudit', '97.1%', '100%', '54'],
        ['preRetrievalDecision', '93.5%', '100%', '46'],
        ['confidenceEvolution', '87.2%', '100%', '16'],
        ['总计', '-', '-', '~654+'],
    ]
    for row_idx, row_data in enumerate(tests):
        row_cells = table4.rows[row_idx].cells
        for col_idx, cell_data in enumerate(row_data):
            row_cells[col_idx].text = cell_data
            if row_idx == 0:
                row_cells[col_idx].paragraphs[0].runs[0].bold = True
                row_cells[col_idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(row_cells[col_idx], '4472C4')
    
    doc.add_paragraph()
    
    # ==================== 9. 项目文档 ====================
    doc.add_heading('9. 项目文档', 1)
    
    table5 = doc.add_table(rows=9, cols=2)
    table5.style = 'Table Grid'
    docs = [
        ['文档', '说明'],
        ['README.md', '项目主文档'],
        ['ARCHITECTURE.md', '系统架构设计'],
        ['INTERFACE_DESIGN.md', '接口设计'],
        ['CLAUDE.md', 'AI 编程指南'],
        ['TESTING.md', '全量测试文档'],
        ['CHANGELOG.md', '变更日志'],
        ['docs/ARCHITECTURE_AIRI_STAGE.md', 'AIRI 舞台架构'],
        ['docs/PRODUCT_SPEC_AIRI_STAGE.md', 'AIRI 产品规格'],
    ]
    for row_idx, row_data in enumerate(docs):
        row_cells = table5.rows[row_idx].cells
        for col_idx, cell_data in enumerate(row_data):
            row_cells[col_idx].text = cell_data
            if row_idx == 0:
                row_cells[col_idx].paragraphs[0].runs[0].bold = True
                row_cells[col_idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(row_cells[col_idx], '4472C4')
    
    doc.add_paragraph()
    
    # ==================== 10. 开发路线 ====================
    doc.add_heading('10. 开发路线（已完成 8 轮迭代）', 1)
    
    milestones = [
        ('第1轮', '个性引擎 + 情感渲染 + 记忆系统整合'),
        ('第2轮', 'PostgreSQL 迁移 + 四层过滤管道 + 自进化闭环'),
        ('第3轮', 'Agent Card 动态发现 + 并行执行引擎 + 委托协议'),
        ('第4轮', '主动记忆引擎（行为检测 + 意图预测 + 预取缓存）'),
        ('第5轮', 'Prompt Caching + Fork 子代理 + DreamGatekeeper'),
        ('第6轮', '记忆技能化改造（Agent 主动调度）'),
        ('第7轮', '记忆系统优化（Embedding + 智能检索 + 质量门控）'),
        ('第8轮', 'AIRI 前端角色舞台集成（Live2D + 事件驱动 + 状态机）'),
    ]
    
    for milestone, desc in milestones:
        p = doc.add_paragraph(f'{milestone}：{desc}')
        p.paragraph_format.left_indent = Inches(0.3)
    
    doc.add_paragraph()
    
    # ==================== 11. 待完成功能 ====================
    doc.add_heading('11. 待完成功能', 1)
    
    todos = [
        '• 闭合自进化反馈回路（classifyNode 消费工具效用分数）',
        '• Agent Card 的 llmConfig 消费端实现',
        '• 迁移到 pgvector 扩展，将向量检索下沉到数据库层',
        '• 引入 Apache AGE 图记忆',
        '• AIRI Bridge 流式输出与前端舞台深度结合',
        '• 前端 UI 适配个性切换和情感渲染展示',
    ]
    
    for todo in todos:
        p = doc.add_paragraph(todo)
        p.paragraph_format.left_indent = Inches(0.3)
    
    # ==================== 12. 总结 ====================
    doc.add_heading('12. 总结', 1)
    
    doc.add_paragraph(
        'SmartAgent4 是一款功能完善、架构清晰的多模态 AI 助手项目。'
        '项目代码总量超过 55,000 行，包含 315 个 TypeScript 文件，'
        '涵盖了从前端 UI 到后端 AI 编排的完整技术栈。'
    )
    
    doc.add_paragraph(
        '项目已完成 8 轮迭代，在多智能体协同、三层记忆系统、AIRI 舞台渲染等核心功能上'
        '具有较强的技术领先性，同时保持了良好的代码质量和测试覆盖率。'
    )
    
    # 保存文档
    import os
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(script_dir)
    output_path = os.path.join(project_dir, 'docs', 'SmartAgent4_项目分析报告_含图表.docx')
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f'\n文档已生成: {output_path}')
    
    # 清理临时图片
    for img in [arch_img, flow_img, agent_img, memory_img]:
        if os.path.exists(img):
            os.remove(img)
    
    return output_path


if __name__ == '__main__':
    create_document()
